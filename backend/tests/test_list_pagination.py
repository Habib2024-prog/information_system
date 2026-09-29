from datetime import datetime, timedelta, timezone
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.audit_log import AuditLog
from app.models.department import Department
from app.models.employee import Employee
from app.models.employee_department import EmployeeDepartment
from app.models.school import School
from app.models.scientific_member import ScientificMember
from app.models.user import User
from app.repositories.user_repository import UserRepository
from tests.test_employees import employee_payload
from tests.test_schools import _school_payload


ENDPOINTS = {
    "employees": "/api/employees",
    "members": "/api/scientific-members",
    "schools": "/api/schools",
    "users": "/api/users",
    "audit": "/api/audit-logs",
}


@pytest.fixture
def admin_client(client: TestClient, authenticated_user: User, db_session: Session) -> TestClient:
    authenticated_user.role_code = "admin"
    db_session.commit()
    return client


def seed_records(db: Session, kind: str, count: int, actor: User) -> list:
    """Seed isolated test records directly, without adding unrelated CRUD audit rows."""
    records = []
    department = None
    if kind in {"employees", "members"}:
        department = Department(code="science")
        db.add(department)
        db.flush()
    for index in range(count - (1 if kind == "users" else 0)):
        name = "Matching" if index < 22 else "Other"
        if kind == "employees":
            values = employee_payload(name=name, father_name=f"Father {index}")
            values.pop("department_ids")
            record = Employee(**values)
        elif kind == "members":
            record = ScientificMember(name=name, surname=f"Surname {index}", father_name="Father",
                                      phone_number="0700000000", academic_rank="استاد",
                                      department_id=department.id, notes="Long details remain available")
        elif kind == "schools":
            values = _school_payload(school_name=name, school_code=f"SCH-{index:03}")
            values.pop("grade_statistics")
            record = School(**values)
        elif kind == "users":
            record = User(username=f"user-{index:03}", full_name=name, role_code="user",
                          password_hash=actor.password_hash, is_active=True)
        else:
            record = AuditLog(user_id=actor.id, action="CREATE" if index < 22 else "UPDATE",
                              entity_type="employee", entity_id=index + 1, description="فعالیت آزمایشی",
                              created_at=datetime.now(timezone.utc) + timedelta(seconds=index))
        db.add(record)
        db.flush()
        if kind == "employees":
            db.add(EmployeeDepartment(employee_id=record.id, department_id=department.id))
        records.append(record)
    db.commit()
    return records


@pytest.mark.parametrize("kind,total", [
    (kind, total) for kind in ENDPOINTS for total in (0, 1, 9, 10, 11, 23)
    if not (kind == "users" and total == 0)
])
def test_default_ten_and_all_page_boundaries(admin_client, db_session, authenticated_user, kind, total):
    seed_records(db_session, kind, total, authenticated_user)
    endpoint = ENDPOINTS[kind]
    first = admin_client.get(endpoint)
    assert first.status_code == 200
    body = first.json()
    assert set(body) == {"items", "total", "page", "page_size"}
    assert (body["page"], body["page_size"], body["total"]) == (1, 10, total)
    assert len(body["items"]) == min(10, total)
    ids = []
    for page in range(1, max(1, (total + 9) // 10) + 1):
        result = admin_client.get(endpoint, params={"page": page}).json()
        assert result["page"] == page and result["page_size"] == 10
        assert result["total"] == total
        assert len(result["items"]) == min(10, max(0, total - (page - 1) * 10))
        ids.extend(item["id"] for item in result["items"])
    assert len(ids) == len(set(ids)) == total
    assert ids == sorted(ids, reverse=kind == "audit")
    beyond = admin_client.get(endpoint, params={"page": max(1, (total + 9) // 10) + 1}).json()
    assert beyond["items"] == [] and beyond["total"] == total


def test_empty_user_repository(db_session):
    # An authenticated user-list API necessarily has at least the administrator.
    assert UserRepository().list_users(db_session, offset=0, limit=10) == ([], 0)


@pytest.mark.parametrize("kind", ["employees", "members", "schools", "audit"])
def test_filters_count_complete_dataset_before_pagination(admin_client, db_session, authenticated_user, kind):
    records = seed_records(db_session, kind, 27, authenticated_user)
    params = {"page": 3, "page_size": 10}
    if kind == "employees":
        department_id = db_session.scalar(select(EmployeeDepartment.department_id).where(EmployeeDepartment.employee_id == records[0].id))
        params.update(search="Matching", job_title_code="teacher", department_id=department_id)
    elif kind == "members":
        params.update(search="Matching", academic_rank="استاد", department_id=records[0].department_id)
    elif kind == "schools":
        params.update(search="Matching", school_type_code="high_school", gender_type_code="mixed")
    else:
        params.update(action="CREATE", entity_type="employee", user_id=authenticated_user.id)
    result = admin_client.get(ENDPOINTS[kind], params=params)
    assert result.status_code == 200
    body = result.json()
    assert body["total"] == 22 and len(body["items"]) == 2
    if kind == "employees":
        assert all(item["subjects_taught"] and item["departments"] for item in body["items"])
    if kind == "members":
        assert all(item["notes"] and item["department"] for item in body["items"])


@pytest.mark.parametrize("kind,sort_by", [("employees", "name"), ("members", "name"), ("schools", "school_name")])
@pytest.mark.parametrize("sort_order", ["asc", "desc"])
def test_equal_sort_values_have_stable_unique_order(admin_client, db_session, authenticated_user, kind, sort_by, sort_order):
    records = seed_records(db_session, kind, 21, authenticated_user)
    ids = []
    for page in (1, 2, 3):
        body = admin_client.get(ENDPOINTS[kind], params={"page": page, "sort_by": sort_by, "sort_order": sort_order}).json()
        ids.extend(item["id"] for item in body["items"])
    assert ids == sorted((record.id for record in records), reverse=sort_order == "desc")


@pytest.mark.parametrize("kind", ["employees", "members", "schools", "users"])
def test_deletion_of_last_row_returns_updated_total_and_valid_previous_page(admin_client, db_session, authenticated_user, kind):
    seed_records(db_session, kind, 21, authenticated_user)
    endpoint = ENDPOINTS[kind]
    last = admin_client.get(endpoint, params={"page": 3}).json()
    assert len(last["items"]) == 1
    assert admin_client.delete(f"{endpoint}/{last['items'][0]['id']}").status_code == 204
    invalid = admin_client.get(endpoint, params={"page": 3}).json()
    assert invalid["items"] == [] and invalid["total"] == 20
    previous = admin_client.get(endpoint, params={"page": 2}).json()
    assert len(previous["items"]) == 10 and previous["total"] == 20


@pytest.mark.parametrize("kind", ["employees", "members", "schools"])
def test_filtered_export_not_limited_to_visible_page(admin_client, db_session, authenticated_user, kind):
    records = seed_records(db_session, kind, 27, authenticated_user)
    endpoint = f"{ENDPOINTS[kind]}/export"
    params = {"search": "Matching"}
    if kind == "employees":
        department_id = db_session.scalar(select(EmployeeDepartment.department_id).where(EmployeeDepartment.employee_id == records[0].id))
        endpoint = f"/api/departments/{department_id}/employees/export"
    response = admin_client.get(endpoint, params=params)
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.worksheets[0].max_row == 23  # header + all 22 matches


@pytest.mark.parametrize("endpoint", ENDPOINTS.values())
@pytest.mark.parametrize("params", [{"page": 0}, {"page_size": 0}, {"page_size": 101}])
def test_invalid_pagination_rejected(admin_client, endpoint, params):
    assert admin_client.get(endpoint, params=params).status_code == 422
