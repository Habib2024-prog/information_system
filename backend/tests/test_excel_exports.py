from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.common.department_labels import DEPARTMENT_DISPLAY_LABELS
from app.common.employee_codes import TEACHER_JOB_TITLE_CODE
from app.models.department import Department
from app.services.department_service import DepartmentService


XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _employee_payload(**overrides: object) -> dict[str, object]:
    payload: dict[str, object] = {
        "name": "احمد",
        "father_name": "کریم",
        "grandfather_name": "رحیم",
        "school_workplace": "مکتب مرکزی",
        "city_district": "کابل",
        "phone_number": "0700000000",
        "field_of_study": "ریاضی",
        "education_level": "لیسانس",
        "subjects_taught": "ریاضی و فزیک",
        "job_title_code": TEACHER_JOB_TITLE_CODE,
        "teaching_experience": 5,
        "grade_post": 4,
        "step": 2,
        "successful_evaluation": "بلی",
        "field_match_code": "in_field",
        "notes": "ملاحظات طولانی کارمند",
        "department_ids": [],
    }
    payload.update(overrides)
    return payload


def _school_payload(**overrides: object) -> dict[str, object]:
    payload: dict[str, object] = {
        "school_name": "لیسه استقلال",
        "school_head_phone": "0700000000",
        "school_type_code": "high_school",
        "gender_type_code": "mixed",
        "school_code": "SCH-001",
        "school_formation": "رسمی",
        "senior_teacher_count": 2,
        "male_teacher_count": 10,
        "female_teacher_count": 8,
        "incoming_service_teacher_count": 1,
        "outgoing_service_teacher_count": 0,
        "volunteer_teacher_count": 3,
        "active_class_section_count": 12,
        "school_needs": "نیاز به کتابخانه و صنف اضافی",
        "school_equipment": "کمپیوتر و میزهای آموزشی",
        "grade_statistics": [],
    }
    payload.update(overrides)
    return payload


def _grade(grade_number: int, **overrides: object) -> dict[str, object]:
    statistic: dict[str, object] = {
        "grade_number": grade_number,
        "enrolled_count": 30,
        "present_count": 28,
        "male_count": 15,
        "female_count": 15,
    }
    statistic.update(overrides)
    return statistic


def _worksheet(response):
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(XLSX_MEDIA_TYPE)
    assert "attachment; filename=" in response.headers["content-disposition"]
    return load_workbook(BytesIO(response.content)).active


def _rows_by_header(worksheet) -> list[dict[str, object]]:
    headers = [cell.value for cell in worksheet[1]]
    return [
        dict(zip(headers, row, strict=True))
        for row in worksheet.iter_rows(min_row=2, values_only=True)
    ]


def test_employee_export_includes_complete_details_labels_and_multiple_departments(
    client: TestClient,
    db_session: Session,
) -> None:
    DepartmentService().seed_predefined_departments(db_session)
    departments = {
        department.code: department
        for department in db_session.scalars(select(Department))
    }
    created = client.post(
        "/api/employees",
        json=_employee_payload(
            department_ids=[departments["science"].id, departments["mathematics"].id],
            field_match_code="out_of_field",
        ),
    )
    assert created.status_code == 201

    worksheet = _worksheet(client.get("/api/employees/export"))
    rows = _rows_by_header(worksheet)

    assert worksheet.title == "کارمندان"
    assert worksheet.sheet_view.rightToLeft is True
    assert worksheet.freeze_panes == "A2"
    assert worksheet.auto_filter.ref == "A1:R2"
    assert len(worksheet[1]) == 18
    assert rows == [
        {
            "شماره": created.json()["id"],
            "اسم": "احمد",
            "ولد": "کریم",
            "ولدیت": "رحیم",
            "مکتب / محل وظیفه": "مکتب مرکزی",
            "شهر / ولسوالی": "کابل",
            "شماره تماس": "0700000000",
            "رشته تحصیلی": "ریاضی",
            "درجه تحصیل": "لیسانس",
            "مضامین که تدریس می‌کند": "ریاضی و فزیک",
            "عنوان وظیفه": "معلم",
            "سابقه تدریس": 5,
            "بست": 4,
            "قدم": 2,
            "ارزیابی موفق": "بلی",
            "مطابق رشته": "خلاف رشته",
            "دیپارتمنت مربوطه": "ریاضی، ساینس",
            "ملاحظات": "ملاحظات طولانی کارمند",
        }
    ]
    assert worksheet.cell(row=2, column=7).data_type == "s"


def test_employee_export_uses_filters_ignores_pagination_and_excludes_deleted(
    client: TestClient,
) -> None:
    first = client.post("/api/employees", json=_employee_payload(name="احمد"))
    second = client.post("/api/employees", json=_employee_payload(name="سلیم", phone_number="0700000001"))
    deleted = client.post("/api/employees", json=_employee_payload(name="حذف", phone_number="0700000002"))
    assert first.status_code == second.status_code == deleted.status_code == 201
    assert client.delete(f"/api/employees/{deleted.json()['id']}").status_code == 204

    all_rows = _rows_by_header(
        _worksheet(client.get("/api/employees/export", params={"page": 1, "page_size": 1}))
    )
    filtered_rows = _rows_by_header(
        _worksheet(client.get("/api/employees/export", params={"name": "سلیم"}))
    )

    assert [row["اسم"] for row in all_rows] == ["احمد", "سلیم"]
    assert [row["اسم"] for row in filtered_rows] == ["سلیم"]


def test_department_employee_export_has_only_selected_department_members_and_full_details(
    client: TestClient,
    db_session: Session,
) -> None:
    DepartmentService().seed_predefined_departments(db_session)
    departments = {
        department.code: department
        for department in db_session.scalars(select(Department))
    }
    science_employee = client.post(
        "/api/employees",
        json=_employee_payload(name="ساینس", department_ids=[departments["science"].id]),
    )
    other_employee = client.post(
        "/api/employees",
        json=_employee_payload(
            name="ریاضی",
            phone_number="0700000001",
            department_ids=[departments["mathematics"].id],
        ),
    )
    assert science_employee.status_code == other_employee.status_code == 201

    worksheet = _worksheet(
        client.get(f"/api/departments/{departments['science'].id}/employees/export")
    )
    rows = _rows_by_header(worksheet)

    assert [row["اسم"] for row in rows] == ["ساینس"]
    assert rows[0]["مضامین که تدریس می‌کند"] == "ریاضی و فزیک"
    assert rows[0]["ملاحظات"] == "ملاحظات طولانی کارمند"


def test_school_export_includes_filtered_complete_school_and_grade_statistics(
    client: TestClient,
) -> None:
    first = client.post(
        "/api/schools",
        json=_school_payload(
            grade_statistics=[
                _grade(1, enrolled_count=31, present_count=30),
                _grade(12, enrolled_count=22, present_count=21),
            ]
        ),
    )
    second = client.post(
        "/api/schools",
        json=_school_payload(
            school_name="مکتب دخترانه",
            school_code="SCH-002",
            school_type_code="middle_school",
            gender_type_code="girls",
        ),
    )
    assert first.status_code == second.status_code == 201

    worksheet = _worksheet(
        client.get("/api/schools/export", params={"school_type_code": "high_school"})
    )
    rows = _rows_by_header(worksheet)
    row = rows[0]

    assert worksheet.title == "مکاتب"
    assert len(rows) == 1
    assert row["نام مکتب"] == "لیسه استقلال"
    assert row["نوع مکتب"] == "لیسه"
    assert row["نوع جنسیت"] == "مختلط"
    assert row["شماره تماس آمر مکتب"] == "0700000000"
    assert row["نیازهای مکتب"] == "نیاز به کتابخانه و صنف اضافی"
    assert row["تجهیزات مکتب"] == "کمپیوتر و میزهای آموزشی"
    headers = [cell.value for cell in worksheet[1]]
    grade_headers = [header for header in headers if isinstance(header, str) and header.startswith("صنف ")]
    first_grade_cell = row["صنف ۱"]
    twelfth_grade_cell = row["صنف ۱۲"]

    assert grade_headers == [f"صنف {number}" for number in "۱۲۳۴۵۶۷۸۹"] + ["صنف ۱۰", "صنف ۱۱", "صنف ۱۲"]
    assert len(grade_headers) == 12
    assert headers.count("صنف ۱") == 1
    assert headers.count("صنف ۱۲") == 1
    assert "داخله: 31" in first_grade_cell
    assert "حاضر: 30" in first_grade_cell
    assert "ذکور: 15" in first_grade_cell
    assert "اناث: 15" in first_grade_cell
    assert first_grade_cell.count("\n") == 3
    assert "داخله: 22" in twelfth_grade_cell
    assert row["صنف ۲"] is None
    assert worksheet.cell(row=2, column=headers.index("صنف ۱") + 1).alignment.wrap_text is True
    assert worksheet.row_dimensions[2].height == 60
    assert worksheet.cell(row=2, column=3).data_type == "s"


def test_school_export_excludes_soft_deleted_schools(client: TestClient) -> None:
    active = client.post("/api/schools", json=_school_payload())
    deleted = client.post(
        "/api/schools",
        json=_school_payload(school_code="SCH-002", school_name="مکتب حذف‌شده"),
    )
    assert active.status_code == deleted.status_code == 201
    assert client.delete(f"/api/schools/{deleted.json()['id']}").status_code == 204

    rows = _rows_by_header(_worksheet(client.get("/api/schools/export")))

    assert [row["نام مکتب"] for row in rows] == ["لیسه استقلال"]
