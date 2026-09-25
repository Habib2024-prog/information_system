from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.common.employee_codes import AMIR_JOB_TITLE_CODE, TEACHER_JOB_TITLE_CODE
from app.models.department import Department
from app.services.department_service import DepartmentService


XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _departments(db_session: Session) -> dict[str, Department]:
    DepartmentService().seed_predefined_departments(db_session)
    return {department.code: department for department in db_session.scalars(select(Department))}


def _create_member(client: TestClient, department_id: int, **overrides: object) -> dict[str, object]:
    payload: dict[str, object] = {
        "name": "مریم",
        "surname": "صادقی",
        "father_name": "حکیم",
        "phone_number": "0790000000",
        "academic_rank": "پوهنمل",
        "department_id": department_id,
        "notes": "ملاحظات تفصیلی عضو علمی",
    }
    payload.update(overrides)
    response = client.post("/api/scientific-members", json=payload)
    assert response.status_code == 201
    return response.json()


def _create_employee(client: TestClient, job_title_code: str, **overrides: object) -> dict[str, object]:
    payload: dict[str, object] = {
        "name": "سمیرا",
        "father_name": "کریم",
        "grandfather_name": "رحیم",
        "school_workplace": "لیسه مرکزی",
        "city_district": "کابل",
        "phone_number": "0700000000",
        "field_of_study": "ریاضی",
        "education_level": "لیسانس",
        "subjects_taught": "ریاضی",
        "job_title_code": job_title_code,
        "teaching_experience": 5,
        "grade_post": 4,
        "step": 2,
        "successful_evaluation": "بلی",
        "field_match_code": "in_field",
        "notes": "ملاحظات کارمند",
        "department_ids": [],
    }
    payload.update(overrides)
    response = client.post("/api/employees", json=payload)
    assert response.status_code == 201
    return response.json()


def _create_teacher_observation(
    client: TestClient,
    employee_id: int,
    observer_id: int,
) -> dict[str, object]:
    response = client.post(
        f"/api/employees/{employee_id}/teacher-observations",
        json={
            "observer_scientific_member_id": observer_id,
            "observation_date": "2026-09-10",
            "observed_class": "صنف دهم",
            "subject": "ریاضی",
            "subject_knowledge_score": "1.00",
            "lesson_plan_score": "1.00",
            "classroom_management_score": "1.00",
            "assessment_score": "1.00",
            "professional_learning_score": "1.00",
            "community_engagement_score": "1.00",
            "strengths": "نکات قوت",
            "improvements": "نکات قابل اصلاح",
            "notes": "ملاحظات",
        },
    )
    assert response.status_code == 201
    return response.json()


def _create_amir_observation(
    client: TestClient,
    employee_id: int,
    observer_id: int,
) -> dict[str, object]:
    response = client.post(
        f"/api/employees/{employee_id}/amir-observations",
        json={
            "observer_scientific_member_id": observer_id,
            "observation_date": "2026-09-20",
            "observed_class": "صنف دهم",
            "subject": "رهبری",
            "responsibility_score": "1.00",
            "professional_leadership_score": "1.00",
            "community_relations_score": "1.00",
            "professional_development_score": "1.00",
            "strengths": "نکات قوت",
            "improvements": "نکات قابل اصلاح",
            "notes": "ملاحظات",
        },
    )
    assert response.status_code == 201
    return response.json()


def _worksheet(response):
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(XLSX_MEDIA_TYPE)
    assert 'attachment; filename="scientific_members.xlsx"' in response.headers[
        "content-disposition"
    ]
    return load_workbook(BytesIO(response.content)).active


def _rows_by_header(worksheet) -> list[dict[str, object]]:
    headers = [cell.value for cell in worksheet[1]]
    return [
        dict(zip(headers, values, strict=True))
        for values in worksheet.iter_rows(min_row=2, values_only=True)
    ]


def test_scientific_member_export_has_full_persian_details_and_ignores_pagination(
    client: TestClient,
    db_session: Session,
) -> None:
    departments = _departments(db_session)
    science_member = _create_member(client, departments["science"].id)
    mathematics_member = _create_member(
        client,
        departments["mathematics"].id,
        name="فرید",
        surname="احمدی",
        phone_number="0790000001",
    )

    worksheet = _worksheet(
        client.get("/api/scientific-members/export", params={"page": 1, "page_size": 1})
    )
    headers = [cell.value for cell in worksheet[1]]
    rows = _rows_by_header(worksheet)

    assert worksheet.title == "اعضای علمی"
    assert worksheet.sheet_view.rightToLeft is True
    assert worksheet.freeze_panes == "A2"
    assert worksheet.auto_filter.ref == "A1:I3"
    assert headers == [
        "شماره",
        "اسم",
        "تخلص",
        "ولد",
        "شماره تماس",
        "رتبه علمی",
        "دیپارتمنت",
        "تعداد مشاهدات",
        "ملاحظات",
    ]
    assert [row["شماره"] for row in rows] == [science_member["id"], mathematics_member["id"]]
    assert rows[0]["دیپارتمنت"] == "ساینس"
    assert rows[0]["تعداد مشاهدات"] == 0
    assert rows[0]["ملاحظات"] == "ملاحظات تفصیلی عضو علمی"
    assert worksheet.cell(row=2, column=5).data_type == "s"
    assert worksheet.cell(row=2, column=9).alignment.wrap_text is True


def test_scientific_member_export_filters_and_excludes_soft_deleted_members(
    client: TestClient,
    db_session: Session,
) -> None:
    departments = _departments(db_session)
    expected = _create_member(client, departments["science"].id, name="لیلا")
    deleted = _create_member(client, departments["science"].id, name="حذف")
    _create_member(client, departments["mathematics"].id, name="لیلا", surname="دیگر")
    assert client.delete(f"/api/scientific-members/{deleted['id']}").status_code == 204

    worksheet = _worksheet(
        client.get(
            "/api/scientific-members/export",
            params={"name": "لیلا", "department_id": departments["science"].id},
        )
    )
    rows = _rows_by_header(worksheet)

    assert [row["شماره"] for row in rows] == [expected["id"]]


def test_scientific_member_export_sums_active_observations_and_excludes_deleted(
    client: TestClient,
    db_session: Session,
) -> None:
    member = _create_member(client, _departments(db_session)["science"].id)
    teacher = _create_employee(client, TEACHER_JOB_TITLE_CODE)
    amir = _create_employee(client, AMIR_JOB_TITLE_CODE)
    teacher_observation = _create_teacher_observation(client, teacher["id"], member["id"])
    _create_amir_observation(client, amir["id"], member["id"])

    rows_before_delete = _rows_by_header(
        _worksheet(client.get("/api/scientific-members/export"))
    )
    assert rows_before_delete[0]["تعداد مشاهدات"] == 2

    assert (
        client.delete(
            f"/api/employees/{teacher['id']}/teacher-observations/{teacher_observation['id']}"
        ).status_code
        == 204
    )
    rows_after_delete = _rows_by_header(
        _worksheet(client.get("/api/scientific-members/export"))
    )

    assert rows_after_delete[0]["تعداد مشاهدات"] == 1
