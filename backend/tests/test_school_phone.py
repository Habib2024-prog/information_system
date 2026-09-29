from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy.orm import Session

from app.models.school import School
from app.schemas.school import SchoolCreate, SchoolUpdate
from tests.test_schools import _school_payload


@pytest.mark.parametrize("schema", [SchoolCreate, SchoolUpdate])
@pytest.mark.parametrize("phone", [
    "0701234567", "0791234567", " 0701234567 ", "+93701234567",
    "+93 70 123 4567", "0093 70 123 4567", "(070) 123-4567",
    "+93 (0) 701-234-567", "070.123.4567", "۰۷۰۱۲۳۴۵۶۷", "٠٧٠١٢٣٤٥٦٧",
])
def test_school_phone_preserves_text_and_formatting(schema, phone: str) -> None:
    record = schema(**_school_payload(school_head_phone=phone))
    assert record.school_head_phone == phone.strip()
    assert isinstance(record.school_head_phone, str)


@pytest.mark.parametrize("schema", [SchoolCreate, SchoolUpdate])
@pytest.mark.parametrize("phone", [None, "", "   "])
def test_school_phone_is_optional(schema, phone: str | None) -> None:
    assert schema(**_school_payload(school_head_phone=phone)).school_head_phone is None


@pytest.mark.parametrize("schema", [SchoolCreate, SchoolUpdate])
@pytest.mark.parametrize("phone", [
    "abc", "070abc4567", "شماره", "123", "1234567890123456", "++93701234567",
    "070--1234567", "070...1234567", "070/1234567", "(0701234567",
    "0701234567)", "070()1234567", "0701234567-", "070\n1234567",
])
def test_school_phone_rejects_letters_and_malformed_values(schema, phone: str) -> None:
    with pytest.raises(ValidationError, match="شماره تماس معتبر نیست"):
        schema(**_school_payload(school_head_phone=phone))


def test_school_create_edit_phone_and_validation_feedback(client: TestClient) -> None:
    payload = _school_payload(school_head_phone=" 0701234567 ")
    created = client.post("/api/schools", json=payload)
    assert created.status_code == 201
    school_id = created.json()["id"]
    assert created.json()["school_head_phone"] == "0701234567"
    assert client.get(f"/api/schools/{school_id}").json()["school_head_phone"] == "0701234567"
    for invalid in ["079abc4567", "070--1234567"]:
        response = client.put(f"/api/schools/{school_id}", json={**payload, "school_head_phone": invalid})
        assert response.status_code == 422
        assert any(error["loc"][-1] == "school_head_phone" and "شماره تماس معتبر نیست" in error["msg"] for error in response.json()["detail"])
        assert client.get(f"/api/schools/{school_id}").json()["school_head_phone"] == "0701234567"
    for phone, expected in [("+93 (0) 701-234-567", "+93 (0) 701-234-567"), ("", None)]:
        updated = client.put(f"/api/schools/{school_id}", json={**payload, "school_head_phone": phone})
        assert updated.status_code == 200
        assert updated.json()["school_head_phone"] == expected
    assert client.post("/api/schools", json=_school_payload(school_code="INVALID", school_head_phone="letters")).status_code == 422


def test_existing_formatted_phone_remains_editable(client: TestClient, db_session: Session) -> None:
    # Simulate a previously stored value without rewriting or migrating it.
    school = School(**{key: value for key, value in _school_payload(school_head_phone="(070) 123-4567").items() if key != "grade_statistics"})
    db_session.add(school)
    db_session.commit()
    assert client.get(f"/api/schools/{school.id}").json()["school_head_phone"] == "(070) 123-4567"
    response = client.put(f"/api/schools/{school.id}", json=_school_payload(school_name="نام جدید", school_head_phone="(070) 123-4567"))
    assert response.status_code == 200
    assert response.json()["school_head_phone"] == "(070) 123-4567"


def test_school_export_labels_and_phone_text(client: TestClient) -> None:
    assert client.post("/api/schools", json=_school_payload(school_head_phone="0701234567")).status_code == 201
    exported = client.get("/api/schools/export")
    assert exported.status_code == 200
    sheet = load_workbook(BytesIO(exported.content))["مکاتب"]
    headers = [cell.value for cell in sheet[1]]
    assert "خدمتی ورودی" in headers
    assert "خدمتی خروجی" in headers
    assert not any("خدماتی" in str(header) for header in headers)
    phone = sheet.cell(2, headers.index("شماره تماس آمر مکتب") + 1)
    assert phone.value == "0701234567"
    assert phone.data_type == "s"
