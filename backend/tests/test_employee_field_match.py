from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import CheckConstraint, Text
from sqlalchemy.orm import Session

from app.common.employee_codes import get_field_match_display_label
from app.models.employee import Employee
from app.schemas.employee import EmployeeCreate, EmployeeUpdate
from tests.test_employees import employee_payload


STATUSES = [
    ("in_field", "مطابق رشته"),
    ("out_of_field", "مخالف رشته"),
    ("non_professional", "غیر مسلکی"),
]


@pytest.mark.parametrize(("code", "label"), STATUSES)
def test_field_match_create_edit_detail_and_export(
    client: TestClient, db_session: Session, code: str, label: str
) -> None:
    payload = employee_payload(field_match_code=code)
    assert EmployeeCreate.model_validate(payload).field_match_code == code
    assert EmployeeUpdate.model_validate(payload).field_match_code == code
    assert get_field_match_display_label(code) == label

    created = client.post("/api/employees", json=payload)
    assert created.status_code == 201
    employee_id = created.json()["id"]
    assert created.json()["field_match_code"] == code

    # Editing another field must not rewrite the stored status.
    updated = client.put(
        f"/api/employees/{employee_id}", json={**payload, "name": "نام جدید"}
    )
    assert updated.status_code == 200
    assert updated.json()["field_match_code"] == code
    detail = client.get(f"/api/employees/{employee_id}")
    assert detail.status_code == 200
    assert detail.json()["field_match_code"] == code
    assert db_session.get(Employee, employee_id).field_match_code == code

    exported = client.get("/api/employees/export", params={"field_match_code": code})
    assert exported.status_code == 200
    worksheet = load_workbook(BytesIO(exported.content)).active
    headers = [cell.value for cell in worksheet[1]]
    row = dict(zip(headers, next(worksheet.iter_rows(min_row=2, values_only=True)), strict=True))
    assert row["مطابق رشته"] == label
    assert row["اسم"] == "نام جدید"
    assert row["شماره تماس"] == payload["phone_number"]


def test_can_update_existing_status_to_non_professional(client: TestClient) -> None:
    created = client.post("/api/employees", json=employee_payload(field_match_code="out_of_field"))
    assert created.status_code == 201
    response = client.put(
        f"/api/employees/{created.json()['id']}",
        json=employee_payload(field_match_code="non_professional"),
    )
    assert response.status_code == 200
    assert response.json()["field_match_code"] == "non_professional"


def test_non_professional_filter_returns_complete_matching_records_and_export(
    client: TestClient,
) -> None:
    for code, _ in STATUSES:
        assert client.post(
            "/api/employees", json=employee_payload(name=code, field_match_code=code)
        ).status_code == 201

    filters = {"field_match_code": "non_professional", "job_title_code": "teacher"}
    response = client.get("/api/employees", params=filters)
    assert response.status_code == 200
    assert response.json()["total"] == 1
    [employee] = response.json()["items"]
    assert employee["field_match_code"] == "non_professional"
    assert employee["notes"] == "Employee note"
    assert employee["subjects_taught"] == "Mathematics and Physics"

    exported = client.get("/api/employees/export", params=filters)
    assert exported.status_code == 200
    worksheet = load_workbook(BytesIO(exported.content)).active
    assert worksheet.max_row == 2
    headers = [cell.value for cell in worksheet[1]]
    assert worksheet.cell(2, headers.index("مطابق رشته") + 1).value == "غیر مسلکی"


def test_legacy_text_status_is_not_rejected_or_rewritten(client: TestClient) -> None:
    payload = employee_payload(field_match_code="legacy_status")
    created = client.post("/api/employees", json=payload)
    assert created.status_code == 201
    updated = client.put(
        f"/api/employees/{created.json()['id']}", json={**payload, "notes": "Updated note"}
    )
    assert updated.status_code == 200
    assert updated.json()["field_match_code"] == "legacy_status"


@pytest.mark.parametrize("schema", [EmployeeCreate, EmployeeUpdate])
def test_field_match_is_still_required(schema: type[EmployeeCreate] | type[EmployeeUpdate]) -> None:
    with pytest.raises(ValidationError):
        schema.model_validate(employee_payload(field_match_code=""))


def test_field_match_database_column_remains_unrestricted_text() -> None:
    assert isinstance(Employee.__table__.c.field_match_code.type, Text)
    assert not any(
        isinstance(constraint, CheckConstraint) and "field_match_code" in str(constraint.sqltext)
        for constraint in Employee.__table__.constraints
    )
