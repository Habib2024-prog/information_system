from io import BytesIO
from decimal import Decimal

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.common.employee_codes import AMIR_JOB_TITLE_CODE, TEACHER_JOB_TITLE_CODE
from app.common.amir_competency_levels import format_amir_competency
from app.common.teacher_competency_levels import format_teacher_competency
from tests.test_scientific_member_observation_history import (
    _create_amir_observation,
    _create_employee,
    _create_member,
    _create_teacher_observation,
    _science_department_id,
)


XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _workbook(response):
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(XLSX_MEDIA_TYPE)
    assert "attachment; filename=" in response.headers["content-disposition"]
    return load_workbook(BytesIO(response.content))


def _rows_by_header(worksheet) -> list[dict[str, object]]:
    headers = [cell.value for cell in worksheet[1]]
    return [
        dict(zip(headers, values, strict=True))
        for values in worksheet.iter_rows(min_row=2, values_only=True)
    ]


def test_teacher_observation_export_has_joined_details_scores_filters_and_no_pagination(
    client: TestClient,
    db_session: Session,
) -> None:
    member = _create_member(client, _science_department_id(db_session))
    teacher = _create_employee(
        client,
        TEACHER_JOB_TITLE_CODE,
        name="معلم اول",
        father_name="احمد",
        school_workplace="لیسه مرکزی",
    )
    second_teacher = _create_employee(
        client,
        TEACHER_JOB_TITLE_CODE,
        name="معلم دوم",
        father_name="ناصر",
        phone_number="0700000001",
    )
    first = _create_teacher_observation(
        client,
        teacher["id"],
        member["id"],
        subject_knowledge_score="0.50",
        lesson_plan_score="1.20",
        classroom_management_score="2.00",
        assessment_score="2.80",
    )
    second = _create_teacher_observation(
        client,
        second_teacher["id"],
        member["id"],
        subject="فزیک",
    )
    deleted = _create_teacher_observation(
        client,
        teacher["id"],
        member["id"],
        subject="حذف‌شده",
    )
    assert (
        client.delete(
            f"/api/employees/{teacher['id']}/teacher-observations/{deleted['id']}"
        ).status_code
        == 204
    )

    workbook = _workbook(
        client.get("/api/teacher-observations/export", params={"page": 1, "page_size": 1})
    )
    worksheet = workbook.active
    rows = _rows_by_header(worksheet)

    assert worksheet.title == "مشاهدات معلمین"
    assert [row["شماره"] for row in rows] == [first["id"], second["id"]]
    assert rows[0]["اسم کارمند"] == "معلم اول"
    assert rows[0]["ولد"] == "احمد"
    assert rows[0]["محل وظیفه / مکتب"] == "لیسه مرکزی"
    assert rows[0]["عنوان وظیفه"] == "معلم"
    assert rows[0]["مشاهده‌کننده"] == "مریم صادقی"
    assert [
        rows[0][header]
        for header in (
            "دانش مضمونی",
            "پلان درسی",
            "مدیریت صنف",
            "ارزیابی",
            "آموزش‌های مسلکی",
            "ارتباط با اجتماع",
        )
    ] == [
        "0.50 - قابلیت مشاهده نشد",
        "1.20 - نیازمند بهبود",
        "2.00 - دارای قابلیت",
        "2.80 - تسلط بر قابلیت",
        "1.00 - نیازمند بهبود",
        "1.00 - نیازمند بهبود",
    ]
    assert rows[0]["مجموع نمره"] == 8.5
    assert rows[0]["نتیجه نهایی"] == "نیازمند بهبود"
    assert worksheet.cell(row=2, column=18).alignment.wrap_text is True
    assert [cell.value for cell in worksheet[1]][9:15] == [
        "دانش مضمونی",
        "پلان درسی",
        "مدیریت صنف",
        "ارزیابی",
        "آموزش‌های مسلکی",
        "ارتباط با اجتماع",
    ]

    filtered_rows = _rows_by_header(
        _workbook(
            client.get(
                "/api/teacher-observations/export",
                params={"employee_id": second_teacher["id"], "subject": "فزیک"},
            )
        ).active
    )
    assert [row["شماره"] for row in filtered_rows] == [second["id"]]


def test_amir_observation_export_has_separate_competencies_translated_result_and_filters(
    client: TestClient,
    db_session: Session,
) -> None:
    member = _create_member(client, _science_department_id(db_session))
    amir = _create_employee(
        client,
        AMIR_JOB_TITLE_CODE,
        name="آمر اول",
        father_name="حکیم",
    )
    other_amir = _create_employee(
        client,
        AMIR_JOB_TITLE_CODE,
        name="آمر دوم",
        phone_number="0700000002",
    )
    first = _create_amir_observation(
        client,
        amir["id"],
        member["id"],
        responsibility_score="0.50",
        professional_leadership_score="1.20",
        community_relations_score="2.00",
        professional_development_score="2.80",
    )
    second = _create_amir_observation(
        client,
        other_amir["id"],
        member["id"],
        subject="مدیریت",
    )
    deleted = _create_amir_observation(client, amir["id"], member["id"], subject="حذف‌شده")
    assert (
        client.delete(f"/api/employees/{amir['id']}/amir-observations/{deleted['id']}").status_code
        == 204
    )

    worksheet = _workbook(client.get("/api/amir-observations/export")).active
    rows = _rows_by_header(worksheet)

    assert worksheet.title == "مشاهدات آمر و سرمعلم"
    assert [row["شماره"] for row in rows] == [first["id"], second["id"]]
    assert rows[0]["اسم کارمند"] == "آمر اول"
    assert rows[0]["مشاهده‌کننده"] == "مریم صادقی"
    assert [
        rows[0][header]
        for header in (
            "مسوولیت پذیری",
            "رهبری مسلکی",
            "روابط با جامعه",
            "انکشاف مسلکی",
        )
    ] == [
        "0.50 - غیر قابل ارزیابی",
        "1.20 - نیازمند بهبود",
        "2.00 - دارای قابلیت",
        "2.80 - تسلط بر قابلیت",
    ]
    assert rows[0]["نتیجه نهایی"] == "قابلیت بکارگیری"
    assert [cell.value for cell in worksheet[1]][9:13] == [
        "مسوولیت پذیری",
        "رهبری مسلکی",
        "روابط با جامعه",
        "انکشاف مسلکی",
    ]

    filtered_rows = _rows_by_header(
        _workbook(
            client.get(
                "/api/amir-observations/export",
                params={"employee_id": other_amir["id"], "subject": "مدیریت"},
            )
        ).active
    )
    assert [row["شماره"] for row in filtered_rows] == [second["id"]]


def test_scientific_member_observation_history_export_uses_two_separate_sheets(
    client: TestClient,
    db_session: Session,
) -> None:
    department_id = _science_department_id(db_session)
    member = _create_member(client, department_id)
    other_member = _create_member(client, department_id, name="فرید", surname="احمدی")
    teacher = _create_employee(client, TEACHER_JOB_TITLE_CODE)
    amir = _create_employee(client, AMIR_JOB_TITLE_CODE)
    teacher_observation = _create_teacher_observation(client, teacher["id"], member["id"])
    amir_observation = _create_amir_observation(client, amir["id"], member["id"])
    _create_teacher_observation(client, teacher["id"], other_member["id"], subject="مشاهده دیگر")

    workbook = _workbook(
        client.get(f"/api/scientific-members/{member['id']}/observations/export")
    )
    teacher_sheet = workbook["مشاهدات معلمین"]
    amir_sheet = workbook["مشاهدات آمر و سرمعلم"]

    assert workbook.sheetnames == ["مشاهدات معلمین", "مشاهدات آمر و سرمعلم"]
    assert [row["شماره"] for row in _rows_by_header(teacher_sheet)] == [teacher_observation["id"]]
    assert [row["شماره"] for row in _rows_by_header(amir_sheet)] == [amir_observation["id"]]
    assert "دانش مضمونی" in [cell.value for cell in teacher_sheet[1]]
    assert "مسوولیت پذیری" in [cell.value for cell in amir_sheet[1]]


def test_competency_level_helpers_leave_undefined_decimal_gaps_blank() -> None:
    assert format_teacher_competency(Decimal("2.10")) == "2.10 - دارای قابلیت"
    assert format_amir_competency(Decimal("2.10")) == "2.10 - دارای قابلیت"
    assert format_teacher_competency(Decimal("0.755")) == "0.755"
    assert format_teacher_competency(Decimal("1.55")) == "1.55"
    assert format_teacher_competency(Decimal("2.255")) == "2.255"
    assert format_amir_competency(Decimal("0.755")) == "0.755"
    assert format_amir_competency(Decimal("1.55")) == "1.55"
    assert format_amir_competency(Decimal("2.255")) == "2.255"
