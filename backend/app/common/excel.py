from collections.abc import Iterable, Sequence
from datetime import date, datetime
from io import BytesIO

from fastapi import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def build_workbook(
    *,
    worksheet_title: str,
    headers: Sequence[str],
    rows: Iterable[Sequence[object | None]],
    wrap_text_columns: set[int] | None = None,
    text_columns: set[int] | None = None,
) -> Workbook:
    """Build a readable RTL workbook from domain-specific headers and rows.

    Column indexes are one-based, matching Excel's column numbering.
    """
    workbook = Workbook()
    worksheet = _populate_worksheet(
        workbook.active,
        worksheet_title=worksheet_title,
        headers=headers,
        rows=rows,
        wrap_text_columns=wrap_text_columns,
        text_columns=text_columns,
    )
    return workbook


def add_worksheet(
    workbook: Workbook,
    *,
    worksheet_title: str,
    headers: Sequence[str],
    rows: Iterable[Sequence[object | None]],
    wrap_text_columns: set[int] | None = None,
    text_columns: set[int] | None = None,
) -> None:
    _populate_worksheet(
        workbook.create_sheet(),
        worksheet_title=worksheet_title,
        headers=headers,
        rows=rows,
        wrap_text_columns=wrap_text_columns,
        text_columns=text_columns,
    )


def _populate_worksheet(
    worksheet,
    *,
    worksheet_title: str,
    headers: Sequence[str],
    rows: Iterable[Sequence[object | None]],
    wrap_text_columns: set[int] | None,
    text_columns: set[int] | None,
):
    worksheet.title = worksheet_title
    worksheet.sheet_view.rightToLeft = True
    worksheet.freeze_panes = "A2"

    wrap_text_columns = wrap_text_columns or set()
    text_columns = text_columns or set()

    for column_index, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=column_index, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)

    for row_values in rows:
        row_number = worksheet.max_row + 1
        for column_index, value in enumerate(row_values, start=1):
            cell = worksheet.cell(row=row_number, column=column_index, value=value)
            cell.alignment = Alignment(
                horizontal="right",
                vertical="top",
                wrap_text=column_index in wrap_text_columns,
            )
            if column_index in text_columns and value is not None:
                cell.value = str(value)
                cell.number_format = "@"
            elif isinstance(value, (date, datetime)):
                cell.number_format = "yyyy-mm-dd"

    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{worksheet.max_row}"
    _set_column_widths(worksheet, headers, wrap_text_columns)
    return worksheet


def workbook_download_response(workbook: Workbook, filename: str) -> Response:
    output = BytesIO()
    workbook.save(output)
    return Response(
        content=output.getvalue(),
        media_type=XLSX_MEDIA_TYPE,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _set_column_widths(worksheet, headers: Sequence[str], wrap_text_columns: set[int]) -> None:
    for column_index, header in enumerate(headers, start=1):
        values = [header]
        for row in worksheet.iter_rows(
            min_row=2,
            min_col=column_index,
            max_col=column_index,
            values_only=True,
        ):
            if row[0] is not None:
                values.append(str(row[0]))
        longest_value = max(len(value) for value in values)
        maximum_width = 45 if column_index in wrap_text_columns else 30
        worksheet.column_dimensions[get_column_letter(column_index)].width = min(
            max(longest_value + 2, 12),
            maximum_width,
        )
